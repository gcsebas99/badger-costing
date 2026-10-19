"""Import supplies from the costing template's "Base Datos Productos" sheet (.xlsx).

Columns are located by their header names, so extra/moved columns are fine. Sub-recipe rows (group
"Subrecetas") are skipped: they become recipes in M2. Rows already in the DB (same name + brand,
ignoring case and accents) are skipped, so re-running an import is safe.

The sheet has up to three prices per row (Precio 1–3) but no dates; they're stored with the import
date in column order, so Precio 3 counts as the newest.
"""

from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from sqlmodel import Session, select

from app.models.supply import Supply, SupplyPrice
from app.services.text import fold
from app.services.units import Unit, UnitConversionError, parse_unit

DEFAULT_SHEET = "Base Datos Productos"
SUBRECIPE_GROUP = "subrecetas"
HEADER_SEARCH_ROWS = 10

COLUMNS = {
    "name": "Producto",
    "brand": "Marca comercial",
    "pack_qty": "Formato de compra",
    "pack_unit": "Unidad de medida",
    "group": "Grupo",
    "supplier": "Proveedor",
}
PRICE_COLUMNS = ("Precio 1", "Precio 2", "Precio 3")
FALLBACK_GROUP = "Otros"


class ImportFormatError(ValueError):
    """The workbook doesn't look like the costing template."""


@dataclass
class ImportRow:
    row: int
    name: str
    brand: str | None
    group: str
    supplier: str | None
    pack_qty: Decimal
    pack_unit: Unit
    prices: list[Decimal]


@dataclass
class SkippedRow:
    row: int
    name: str
    reason: str


@dataclass
class ImportPlan:
    rows: list[ImportRow] = field(default_factory=list)
    skipped: list[SkippedRow] = field(default_factory=list)


@dataclass
class ImportResult:
    created: list[ImportRow] = field(default_factory=list)
    existing: list[ImportRow] = field(default_factory=list)


def _text(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _decimal(value: object) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, float):
        value = repr(value)  # shortest exact repr, avoids binary float noise
    try:
        number = Decimal(str(value).strip())
    except InvalidOperation:
        return None
    return number if number.is_finite() else None


def parse_workbook(path: Path, sheet: str = DEFAULT_SHEET) -> ImportPlan:
    workbook = load_workbook(path, data_only=True, read_only=True)
    if sheet not in workbook.sheetnames:
        available = ", ".join(workbook.sheetnames)
        raise ImportFormatError(f"La hoja '{sheet}' no existe. Hojas: {available}")
    rows = list(workbook[sheet].iter_rows(values_only=True))
    workbook.close()

    header_index, columns = _find_header(rows)
    plan = ImportPlan()
    for offset, values in enumerate(rows[header_index + 1 :], start=header_index + 2):

        def cell(header: str, values: tuple = values) -> object:
            index = columns.get(header)
            return values[index] if index is not None and index < len(values) else None

        name = _text(cell(COLUMNS["name"]))
        if name is None:
            continue
        group = _text(cell(COLUMNS["group"])) or FALLBACK_GROUP
        if fold(group) == SUBRECIPE_GROUP:
            plan.skipped.append(SkippedRow(offset, name, "Es una sub-receta (se crea en Recetas)."))
            continue

        pack_qty = _decimal(cell(COLUMNS["pack_qty"]))
        if pack_qty is None or pack_qty <= 0:
            plan.skipped.append(SkippedRow(offset, name, "Falta el formato de compra (cantidad)."))
            continue
        raw_unit = _text(cell(COLUMNS["pack_unit"]))
        if raw_unit is None:
            plan.skipped.append(SkippedRow(offset, name, "Falta la unidad de medida."))
            continue
        try:
            pack_unit = parse_unit(raw_unit)
        except UnitConversionError as exc:
            plan.skipped.append(SkippedRow(offset, name, str(exc)))
            continue

        prices = [
            amount
            for header in PRICE_COLUMNS
            if (amount := _decimal(cell(header))) is not None and amount > 0
        ]
        plan.rows.append(
            ImportRow(
                row=offset,
                name=name,
                brand=_text(cell(COLUMNS["brand"])),
                group=group,
                supplier=_text(cell(COLUMNS["supplier"])),
                pack_qty=pack_qty,
                pack_unit=pack_unit,
                prices=prices,
            )
        )
    return plan


def _find_header(rows: list[tuple]) -> tuple[int, dict[str, int]]:
    for index, values in enumerate(rows[:HEADER_SEARCH_ROWS]):
        names = {_text(v): i for i, v in enumerate(values) if _text(v)}
        if COLUMNS["name"] in names:
            missing = [h for h in (COLUMNS["pack_qty"], COLUMNS["pack_unit"]) if h not in names]
            if missing:
                raise ImportFormatError(f"Faltan columnas: {', '.join(missing)}")
            return index, names
    raise ImportFormatError(f"No se encontró la fila de encabezados (columna '{COLUMNS['name']}').")


def apply_import(session: Session, plan: ImportPlan, purchased_on: date) -> ImportResult:
    existing_keys = {
        (fold(name), fold(brand)) for name, brand in session.exec(select(Supply.name, Supply.brand))
    }
    result = ImportResult()
    for row in plan.rows:
        key = (fold(row.name), fold(row.brand))
        if key in existing_keys:
            result.existing.append(row)
            continue
        existing_keys.add(key)
        supply = Supply(
            name=row.name,
            brand=row.brand,
            group=row.group,
            supplier=row.supplier,
            pack_qty=row.pack_qty,
            pack_unit=row.pack_unit,
        )
        supply.prices = [
            SupplyPrice(amount=amount, purchased_on=purchased_on, supplier=row.supplier)
            for amount in row.prices
        ]
        session.add(supply)
        result.created.append(row)
    session.commit()
    return result
